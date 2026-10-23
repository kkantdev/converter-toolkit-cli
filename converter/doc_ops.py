"""
Document and spreadsheet conversion operations.
"""

from pathlib import Path
from typing import Optional

from rich.console import Console

console = Console()


def resolve_output_path(target_path: Path) -> Path:
    """
    Resolve output path to avoid file collisions.
    
    Args:
        target_path: Desired output path
        
    Returns:
        Safe output path that doesn't overwrite existing files
    """
    if not target_path.exists():
        return target_path

    parent = target_path.parent
    stem = target_path.stem
    suffix = target_path.suffix
    counter = 1

    while True:
        candidate = parent / f"{stem}_cv{counter}{suffix}"
        if not candidate.exists():
            return candidate
        counter += 1


def convert_txt_to_docx(input_path: Path, output_path: Optional[Path] = None) -> Path:
    """
    Convert text file to Word document.
    
    Args:
        input_path: Path to input text file
        output_path: Optional output path (auto-generated if None)
        
    Returns:
        Path to Word document file
        
    Raises:
        Exception: If conversion fails
    """
    try:
        from docx import Document
    except ImportError:
        raise ImportError("python-docx is required for Word document operations. Install with: pip install .[office]")
    
    # Generate output path if not provided
    if output_path is None:
        output_path = input_path.with_suffix('.docx')
    
    # Resolve collision-safe output path
    safe_output_path = resolve_output_path(output_path)
    
    try:
        console.print(f"[cyan]Converting text file {input_path.name} to Word document...[/cyan]")
        
        with console.status("Creating Word document", spinner="dots"):
            # Create new document
            doc = Document()
            
            # Read text file with proper encoding
            try:
                with open(input_path, 'r', encoding='utf-8') as f:
                    content = f.read()
            except UnicodeDecodeError:
                # Fallback to other encodings
                try:
                    with open(input_path, 'r', encoding='latin-1') as f:
                        content = f.read()
                except UnicodeDecodeError:
                    with open(input_path, 'r', encoding='cp1252') as f:
                        content = f.read()
            
            # Split content into paragraphs and add to document
            lines = content.split('\n')
            
            for line in lines:
                line = line.strip()
                if line:  # Only add non-empty lines
                    doc.add_paragraph(line)
                elif len(doc.paragraphs) > 0:  # Add empty paragraph only if there are existing paragraphs
                    doc.add_paragraph()
            
            # Save document
            doc.save(safe_output_path)
        
        console.print(f"[green]✔ Successfully converted text to Word document![/green]")
        console.print(f"[bold]Output saved to:[/bold]\n{safe_output_path}")
        
        return safe_output_path
        
    except Exception as e:
        console.print(f"[red]✖ Error during text to DOCX conversion:[/red] {str(e)}")
        raise


def convert_docx_to_pptx(input_path: Path, output_path: Optional[Path] = None) -> Path:
    """
    Convert Word document to PowerPoint presentation.
    
    Args:
        input_path: Path to input Word document
        output_path: Optional output path (auto-generated if None)
        
    Returns:
        Path to PowerPoint file
        
    Raises:
        Exception: If conversion fails
    """
    try:
        from docx import Document
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        raise ImportError("python-docx and python-pptx are required for document conversion. Install with: pip install .[office]")
    
    # Generate output path if not provided
    if output_path is None:
        output_path = input_path.with_suffix('.pptx')
    
    # Resolve collision-safe output path
    safe_output_path = resolve_output_path(output_path)
    
    try:
        console.print(f"[cyan]Converting Word document {input_path.name} to PowerPoint presentation...[/cyan]")
        
        with console.status("Creating PowerPoint presentation", spinner="dots"):
            # Open Word document
            doc = Document(input_path)
            
            # Create PowerPoint presentation
            prs = Presentation()
            
            # Remove default slide if it exists
            if prs.slides:
                prs.slides._sldIdLst.clear()
            
            # Get slide layout (title and content)
            title_slide_layout = prs.slide_layouts[0]  # Title slide
            content_slide_layout = prs.slide_layouts[1]  # Title and Content
            
            # Track if we've added a title slide
            title_added = False
            current_slide = None
            
            for paragraph in doc.paragraphs:
                text = paragraph.text.strip()
                
                if not text:
                    continue  # Skip empty paragraphs
                
                # First meaningful paragraph becomes title slide
                if not title_added:
                    slide = prs.slides.add_slide(title_slide_layout)
                    title = slide.shapes.title
                    title.text = text
                    title_added = True
                    continue
                
                # Subsequent paragraphs become content slides
                slide = prs.slides.add_slide(content_slide_layout)
                
                # Use first few words as slide title
                words = text.split()[:6]  # First 6 words for title
                slide_title = ' '.join(words)
                if len(words) < len(text.split()):
                    slide_title += '...'
                
                title = slide.shapes.title
                title.text = slide_title
                
                # Add full paragraph text as content
                content = slide.placeholders[1]
                content.text = text
                
                # Format text
                for paragraph in content.text_frame.paragraphs:
                    paragraph.font.size = Pt(18)
            
            # If no content was added, create a single slide with filename
            if not prs.slides:
                slide = prs.slides.add_slide(title_slide_layout)
                title = slide.shapes.title
                title.text = f"Document: {input_path.stem}"
                
                # Add subtitle
                if hasattr(slide, 'placeholders') and len(slide.placeholders) > 1:
                    subtitle = slide.placeholders[1]
                    subtitle.text = "Converted from Word document"
            
            # Save presentation
            prs.save(safe_output_path)
        
        console.print(f"[green]✔ Successfully converted Word document to PowerPoint![/green]")
        console.print(f"[bold]Output saved to:[/bold]\n{safe_output_path}")
        
        return safe_output_path
        
    except Exception as e:
        console.print(f"[red]✖ Error during DOCX to PPTX conversion:[/red] {str(e)}")
        raise


def convert_csv_to_xlsx(input_path: Path, output_path: Optional[Path] = None) -> Path:
    """
    Convert CSV to Excel file.
    
    Args:
        input_path: Path to input CSV file
        output_path: Optional output path (auto-generated if None)
        
    Returns:
        Path to Excel file
        
    Raises:
        Exception: If conversion fails
    """
    try:
        import pandas as pd
    except ImportError:
        raise ImportError("pandas is required for spreadsheet operations. Install with: pip install .[data]")
    
    # Generate output path if not provided
    if output_path is None:
        output_path = input_path.with_suffix('.xlsx')
    
    # Resolve collision-safe output path
    safe_output_path = resolve_output_path(output_path)
    
    try:
        console.print(f"[cyan]Converting CSV {input_path.name} to Excel file...[/cyan]")
        
        with console.status("Processing CSV data", spinner="dots"):
            # Read CSV with multiple encoding attempts
            df = None
            encodings = ['utf-8', 'latin1', 'cp1252', 'iso-8859-1']
            
            for encoding in encodings:
                try:
                    df = pd.read_csv(input_path, encoding=encoding)
                    break
                except UnicodeDecodeError:
                    continue
                except Exception:
                    continue
            
            if df is None:
                raise ValueError("Could not read CSV file with any supported encoding")
            
            # Save to Excel (without index)
            df.to_excel(safe_output_path, index=False, engine='openpyxl')
        
        console.print(f"[green]✔ Successfully converted CSV to Excel![/green]")
        console.print(f"[bold]Rows processed:[/bold] {len(df)}")
        console.print(f"[bold]Columns:[/bold] {len(df.columns)}")
        console.print(f"[bold]Output saved to:[/bold]\n{safe_output_path}")
        
        return safe_output_path
        
    except Exception as e:
        console.print(f"[red]✖ Error during CSV to XLSX conversion:[/red] {str(e)}")
        raise


def convert_xlsx_to_csv(input_path: Path, output_path: Optional[Path] = None) -> Path:
    """
    Convert Excel to CSV file.
    
    Args:
        input_path: Path to input Excel file
        output_path: Optional output path (auto-generated if None)
        
    Returns:
        Path to CSV file
        
    Raises:
        Exception: If conversion fails
    """
    try:
        import pandas as pd
    except ImportError:
        raise ImportError("pandas is required for spreadsheet operations. Install with: pip install .[data]")
    
    # Generate output path if not provided
    if output_path is None:
        output_path = input_path.with_suffix('.csv')
    
    # Resolve collision-safe output path
    safe_output_path = resolve_output_path(output_path)
    
    try:
        console.print(f"[cyan]Converting Excel file {input_path.name} to CSV...[/cyan]")
        
        with console.status("Processing Excel data", spinner="dots"):
            # Read Excel file (first sheet by default)
            df = pd.read_excel(input_path, engine='openpyxl')
            
            # Save to CSV (without index)
            df.to_csv(safe_output_path, index=False, encoding='utf-8')
        
        console.print(f"[green]✔ Successfully converted Excel to CSV![/green]")
        console.print(f"[bold]Rows processed:[/bold] {len(df)}")
        console.print(f"[bold]Columns:[/bold] {len(df.columns)}")
        console.print(f"[bold]Output saved to:[/bold]\n{safe_output_path}")
        
        return safe_output_path
        
    except Exception as e:
        console.print(f"[red]✖ Error during XLSX to CSV conversion:[/red] {str(e)}")
        raise